"""Pull the planning numbers out of GridAid_County_Plans.xlsx into frontend/src/data/logistics.json.

Reads the values Excel last calculated (open, edit and save the workbook in Excel first), finding columns by their
header text so added columns don't break it. Run again whenever the workbook changes:

  python scripts/extract_logistics.py [path/to/GridAid_County_Plans.xlsx]

If GridAid_County_Logistics.xlsx sits next to it, its health-service counts (home oxygen, dialysis, home health,
hospice) and recent-outage rows are merged in too; the logistics page uses them.
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "GridAid_County_Plans.xlsx"
LOGI = XLSX.with_name("GridAid_County_Logistics.xlsx")
SERVICES = {"Home oxygen service": "oxygen", "Dialysis at a center": "dialysis", "Home health": "homeHealth", "Hospice at home": "hospice"}
OUT = ROOT / "frontend" / "src" / "data" / "logistics.json"

GEN_KEYS = {"Hookup only": "hookup", "Permanent generator, diesel": "diesel",
            "Permanent generator, natural gas": "gas", "Solar panels plus battery": "solar"}
INPUTS = {"FEMA share of cost": "fema", "Expected turnout at sites": "turnout", "Power per person, low (watts)": "wattsLow",
          "Power per person, high (watts)": "wattsHigh", "Reach, rural (miles)": "reachRural", "Reach, urban (miles)": "reachUrban",
          "Days in a long outage": "days", "Power per site (kW)": "kwPerSite"}


def find_row(ws, text):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if isinstance(row[0].value, str) and row[0].value.startswith(text):
            return row[0].row
    raise SystemExit(f"{XLSX.name}: no row starting with {text!r} in column A")


def table(ws, header_row):
    """Rows under a header row, as dicts keyed by header text, until the first blank cell in column A."""
    heads = [c.value for c in ws[header_row]]
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if row[0] is None:
            break
        out.append({h: v for h, v in zip(heads, row) if h is not None})
    return out


def starts(row, text):
    """The value under the header that starts with text (the running-cost headers carry the outage length)."""
    return next(v for h, v in row.items() if str(h).startswith(text))


def num(v):
    if not isinstance(v, (int, float)):
        raise SystemExit(f"{XLSX.name}: expected a number, got {v!r}. Open and save the workbook in Excel so values are calculated.")
    return round(v, 2)


def main():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["County Plans"]

    generators = {}
    for g in table(ws, find_row(ws, "Type")):
        key = next((k for name, k in GEN_KEYS.items() if g["Type"].startswith(name)), None)
        if key is None:
            raise SystemExit(f"{XLSX.name}: unknown generator type {g['Type']!r}")
        rent = g["Rental per week, per site ($)"]
        generators[key] = {"name": g["Type"].split(" (")[0], "cost": num(g["Cost per site ($)"]), "fuel": g["Fuel"],
                           "fuelPerDay": num(g["Fuel cost per day, per site ($)"]),
                           "rentPerWeek": num(rent) if isinstance(rent, (int, float)) else None}

    inputs = {}
    for row in table(ws, find_row(ws, "Input")):
        key = next((k for name, k in INPUTS.items() if row["Input"].startswith(name)), None)
        if key:
            inputs[key] = num(row["Value"])
    missing = set(INPUTS.values()) - set(inputs)
    if missing:
        raise SystemExit(f"{XLSX.name}: inputs not found: {', '.join(sorted(missing))}")

    counties = {}
    for c in table(ws, find_row(ws, "County")):
        if not str(c.get("FIPS") or "").strip():
            continue                                              # the total rows
        fips = str(c["FIPS"]).zfill(5)
        counties[fips] = {
            "name": f"{c['County']}, {c['State']}", "type": c["Type"], "pop": num(c["Population (2020)"]),
            "medicare": num(c["Medicare members"]), "rely": num(c["Rely on powered equipment"]),
            "reach": num(c["Reach (mi)"]), "expected": num(c["Expected at sites"]),
            "kw": [num(c["Power needed, low (kW)"]), num(c["Power needed, high (kW)"])],
            "sites": {"recommended": int(c["Recommended sites"]), "high": int(c["Sites, high estimate"])},
            "budget": {
                "recommended": {"hookup": num(c["Hookups only ($)"]), "diesel": num(c["Diesel ($)"]),
                                "gas": num(c["Natural gas ($)"]), "solar": num(c["Solar + battery ($)"])},
                "high": {"hookup": num(c["Hookups only, high ($)"]), "diesel": num(c["Diesel, high ($)"]),
                         "gas": num(c["Natural gas, high ($)"]), "solar": num(c["Solar + battery, high ($)"])},
            },
            "localShareDiesel": num(c["Local share if FEMA pays, diesel ($)"]),
            "running": {"hookup": num(starts(c, "Hookups: rent")), "diesel": num(starts(c, "Diesel fuel")),
                        "gas": num(starts(c, "Natural gas for")), "solar": 0},
        }

    # place names for ZIPs: urban plans name the area, rural plans name the nearest town
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    areas = {}
    if "Urban Site Plans" in wb.sheetnames:
        us = wb["Urban Site Plans"]
        for r in table(us, find_row(us, "County")):
            if r.get("ZIP") and r.get("Area"):
                areas[str(r["ZIP"]).zfill(5)] = r["Area"]
    if "Rural Site Plans" in wb.sheetnames:
        rs = wb["Rural Site Plans"]
        for r in table(rs, find_row(rs, "County")):
            if r.get("ZIP") and isinstance(r.get("Site"), str):
                areas.setdefault(str(r["ZIP"]).zfill(5), r["Site"].replace("Public building near ", ""))

    if LOGI.exists():
        lw = openpyxl.load_workbook(LOGI, data_only=True)
        ls = lw["County Plans"]
        for r in table(ls, find_row(ls, "County")):
            fips = str(r.get("FIPS") or "").zfill(5)
            if fips in counties:
                counties[fips]["services"] = {k: num(r[h]) for h, k in SERVICES.items() if h in r}
        if "Outages" in lw.sheetnames:
            os_ = lw["Outages"]
            for r in table(os_, find_row(os_, "County")):
                for fips, c in counties.items():
                    if c["name"] == r["County"]:
                        c["outage"] = {"date": r["Date"], "what": r["What happened"], "src": r["Source"]}

    OUT.write_text(json.dumps({"source": XLSX.name, "areas": areas, "generators": generators, "inputs": inputs, "counties": counties}, indent=1) + "\n")
    print(f"{OUT.relative_to(ROOT)}: {len(generators)} generator types, {len(counties)} counties")


if __name__ == "__main__":
    main()
